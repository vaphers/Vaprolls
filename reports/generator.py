import os
import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from database.db import Database

logger = logging.getLogger(__name__)

class ReportGenerator:
    def __init__(self, db: Database, audit_id: str, output_dir: str = 'reports_output'):
        self.db = db
        self.audit_id = audit_id
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    async def _fetch_data(self) -> Dict[str, Any]:
        audit = await self.db.get_audit(self.audit_id)
        pages = await self.db.get_pages(self.audit_id)
        issues = await self.db.get_issues(self.audit_id)
        links = await self.db.get_links(self.audit_id)
        images = await self.db.get_images(self.audit_id)
        keywords = await self.db.get_keywords(self.audit_id)
        structured_data = await self.db.get_structured_data(self.audit_id)
        perf_metrics = await self.db.get_performance_metrics(self.audit_id)
        summary = await self.db.get_audit_summary(self.audit_id)
        
        return {
            "audit": audit,
            "pages": pages,
            "issues": issues,
            "links": links,
            "images": images,
            "keywords": keywords,
            "structured_data": structured_data,
            "perf_metrics": perf_metrics,
            "summary": summary
        }

    async def generate_html(self) -> str:
        data = await self._fetch_data()
        domain = data['audit'].get('domain', 'Unknown Domain')
        date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        health_score = data['audit'].get('health_score', 0)
        
        critical_count = len([i for i in data['issues'] if i.get('severity') == 'critical'])
        warning_count = len([i for i in data['issues'] if i.get('severity') == 'warning'])
        info_count = len([i for i in data['issues'] if i.get('severity') == 'info'])
        
        html_content = f"""
        <!DOCTYPE html>
        <html lang="en">
        <head>
            <meta charset="UTF-8">
            <title>SEO Audit Report - {domain}</title>
            <style>
                body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; color: #333; line-height: 1.6; margin: 0; padding: 0; background-color: #f9fafb; }}
                .container {{ max-width: 1200px; margin: 0 auto; padding: 20px; background-color: #fff; box-shadow: 0 0 10px rgba(0,0,0,0.1); }}
                header {{ text-align: center; padding: 40px 0; border-bottom: 2px solid #e5e7eb; }}
                h1 {{ margin: 0; color: #111827; }}
                .health-score {{ font-size: 48px; font-weight: bold; margin: 20px 0; }}
                .score-good {{ color: #10b981; }}
                .score-average {{ color: #f59e0b; }}
                .score-poor {{ color: #ef4444; }}
                .section {{ margin: 40px 0; }}
                .section h2 {{ border-bottom: 1px solid #e5e7eb; padding-bottom: 10px; color: #1f2937; }}
                .card-container {{ display: flex; gap: 20px; flex-wrap: wrap; margin-bottom: 20px; }}
                .card {{ flex: 1; min-width: 200px; padding: 20px; background: #f3f4f6; border-radius: 8px; text-align: center; }}
                .card h3 {{ margin-top: 0; font-size: 14px; text-transform: uppercase; color: #6b7280; }}
                .card p {{ font-size: 24px; font-weight: bold; margin: 0; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                th, td {{ padding: 12px; text-align: left; border-bottom: 1px solid #e5e7eb; }}
                th {{ background-color: #f9fafb; font-weight: 600; color: #374151; }}
                .severity-critical {{ background-color: #fee2e2; color: #b91c1c; padding: 4px 8px; border-radius: 4px; font-size: 12px; font-weight: bold; }}
                .severity-warning {{ background-color: #fef3c7; color: #b45309; padding: 4px 8px; border-radius: 4px; font-size: 12px; font-weight: bold; }}
                .severity-info {{ background-color: #dbeafe; color: #1d4ed8; padding: 4px 8px; border-radius: 4px; font-size: 12px; font-weight: bold; }}
                @media print {{ body {{ background-color: #fff; }} .container {{ box-shadow: none; padding: 0; }} }}
            </style>
        </head>
        <body>
            <div class="container">
                <header>
                    <h1>SEO Audit Report</h1>
                    <p><strong>Domain:</strong> {domain} | <strong>Date:</strong> {date_str}</p>
                    <div class="health-score {'score-good' if health_score >= 80 else 'score-average' if health_score >= 50 else 'score-poor'}">
                        Health Score: {health_score:.1f}/100
                    </div>
                </header>

                <div class="section">
                    <h2>Executive Summary</h2>
                    <div class="card-container">
                        <div class="card">
                            <h3>Total Pages Crawled</h3>
                            <p>{len(data['pages'])}</p>
                        </div>
                        <div class="card">
                            <h3>Critical Issues</h3>
                            <p class="severity-critical">{critical_count}</p>
                        </div>
                        <div class="card">
                            <h3>Warnings</h3>
                            <p class="severity-warning">{warning_count}</p>
                        </div>
                        <div class="card">
                            <h3>Info</h3>
                            <p class="severity-info">{info_count}</p>
                        </div>
                    </div>
                    
                    <h3>Top Priority Issues</h3>
                    <ul>
        """
        
        top_issues = [i for i in data['issues'] if i.get('severity') == 'critical'][:5]
        for issue in top_issues:
            html_content += f"<li><strong>{issue.get('issue_type')}</strong> on {issue.get('url')}: {issue.get('message')}</li>"
            
        html_content += """
                    </ul>
                </div>
                
                <div class="section">
                    <h2>All Issues</h2>
                    <table>
                        <thead>
                            <tr>
                                <th>Severity</th>
                                <th>Type</th>
                                <th>Message</th>
                                <th>URL</th>
                            </tr>
                        </thead>
                        <tbody>
        """
        
        for issue in sorted(data['issues'], key=lambda x: {'critical': 0, 'warning': 1, 'info': 2}.get(x.get('severity', 'info'), 3)):
            severity_class = f"severity-{issue.get('severity', 'info')}"
            html_content += f"""
                            <tr>
                                <td><span class="{severity_class}">{str(issue.get('severity')).upper()}</span></td>
                                <td>{issue.get('issue_type')}</td>
                                <td>{issue.get('message')}</td>
                                <td style="max-width:300px; word-wrap:break-word;">{issue.get('url')}</td>
                            </tr>
            """
            
        html_content += """
                        </tbody>
                    </table>
                </div>
                
                <div class="section">
                    <h2>Page-by-Page Index</h2>
                    <table>
                        <thead>
                            <tr>
                                <th>URL</th>
                                <th>Status Code</th>
                                <th>Title Length</th>
                                <th>Word Count</th>
                            </tr>
                        </thead>
                        <tbody>
        """
        for page in data['pages']:
            title = page.get('title') or ''
            html_content += f"""
                            <tr>
                                <td style="max-width:400px; word-wrap:break-word;">{page.get('url')}</td>
                                <td>{page.get('status_code')}</td>
                                <td>{len(title)}</td>
                                <td>{page.get('word_count')}</td>
                            </tr>
            """
            
        html_content += """
                        </tbody>
                    </table>
                </div>
            </div>
        </body>
        </html>
        """
        
        file_path = self.output_dir / "report.html"
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(html_content)
            
        return str(file_path.absolute())

    async def generate_csv(self) -> str:
        data = await self._fetch_data()
        csv_dir = self.output_dir / "csv_exports"
        csv_dir.mkdir(exist_ok=True)
        
        # issues.csv
        with open(csv_dir / "issues.csv", "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["URL", "Category", "Severity", "Issue Type", "Message", "Recommendation"])
            for issue in data['issues']:
                writer.writerow([
                    issue.get("url"), issue.get("category"), issue.get("severity"),
                    issue.get("issue_type"), issue.get("message"), issue.get("recommendation")
                ])
                
        # pages.csv
        with open(csv_dir / "pages.csv", "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["URL", "Status Code", "Title", "Meta Description", "Word Count", "Response Time", "Issues Count"])
            for page in data['pages']:
                issues_count = len([i for i in data['issues'] if i.get('page_id') == page.get('id')])
                writer.writerow([
                    page.get("url"), page.get("status_code"), page.get("title"),
                    page.get("meta_description"), page.get("word_count"), page.get("response_time_ms"), issues_count
                ])
                
        # links.csv
        with open(csv_dir / "links.csv", "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Source URL", "Target URL", "Type", "Broken", "Status Code", "Anchor Text"])
            for link in data['links']:
                link_type = "Internal" if link.get("is_internal") else "External"
                writer.writerow([
                    link.get("source_url"), link.get("target_url"), link_type,
                    link.get("is_broken"), link.get("status_code"), link.get("anchor_text")
                ])
                
        # images.csv
        with open(csv_dir / "images.csv", "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Page URL", "Image URL", "Alt Text", "Has Dimensions", "Format"])
            for img in data['images']:
                writer.writerow([
                    img.get("page_url"), img.get("src"), img.get("alt_text"),
                    img.get("has_dimensions"), img.get("format")
                ])
                
        # keywords.csv
        with open(csv_dir / "keywords.csv", "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Page URL", "Keyword", "Frequency", "Density", "In Title", "In H1"])
            for kw in data['keywords']:
                writer.writerow([
                    kw.get("page_url"), kw.get("keyword"), kw.get("frequency"),
                    kw.get("density"), kw.get("in_title"), kw.get("in_h1")
                ])
                
        return str(csv_dir.absolute())

    async def generate_excel(self) -> str:
        data = await self._fetch_data()
        wb = openpyxl.Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        def add_sheet(title, headers, rows):
            ws = wb.create_sheet(title=title)
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            for row in rows:
                ws.append(row)
                
        # Issues
        issues_rows = [
            [i.get("url"), i.get("category"), i.get("severity"), i.get("issue_type"), i.get("message"), i.get("recommendation")]
            for i in data['issues']
        ]
        add_sheet("Issues", ["URL", "Category", "Severity", "Issue Type", "Message", "Recommendation"], issues_rows)
        
        # Pages
        pages_rows = []
        for page in data['pages']:
            issues_count = len([i for i in data['issues'] if i.get('page_id') == page.get('id')])
            pages_rows.append([
                page.get("url"), page.get("status_code"), page.get("title"),
                page.get("meta_description"), page.get("word_count"), page.get("response_time_ms"), issues_count
            ])
        add_sheet("Pages", ["URL", "Status Code", "Title", "Meta Description", "Word Count", "Response Time", "Issues Count"], pages_rows)
        
        # Links
        links_rows = [
            [l.get("source_url"), l.get("target_url"), "Internal" if l.get("is_internal") else "External",
             l.get("is_broken"), l.get("status_code"), l.get("anchor_text")]
            for l in data['links']
        ]
        add_sheet("Links", ["Source URL", "Target URL", "Type", "Broken", "Status Code", "Anchor Text"], links_rows)
        
        # Images
        images_rows = [
            [img.get("page_url"), img.get("src"), img.get("alt_text"), img.get("has_dimensions"), img.get("format")]
            for img in data['images']
        ]
        add_sheet("Images", ["Page URL", "Image URL", "Alt Text", "Has Dimensions", "Format"], images_rows)
        
        # Security Headers
        sec_rows = [
            [
                p.get("url"), p.get("status_code"), p.get("hsts_header"), p.get("csp_header"),
                p.get("x_content_type_options"), p.get("x_frame_options"), p.get("referrer_policy")
            ]
            for p in data['pages']
        ]
        add_sheet("Security Headers", ["URL", "Status Code", "HSTS", "CSP", "X-Content-Type-Options", "X-Frame-Options", "Referrer-Policy"], sec_rows)

        # Social Meta (Open Graph & Twitter)
        social_rows = [
            [
                p.get("url"), p.get("og_title"), p.get("og_description"), p.get("og_image"),
                p.get("twitter_card"), p.get("twitter_title"), p.get("twitter_image")
            ]
            for p in data['pages']
        ]
        add_sheet("Social Meta", ["URL", "OG Title", "OG Description", "OG Image", "Twitter Card", "Twitter Title", "Twitter Image"], social_rows)

        # Forms
        forms = await self.db.get_forms(self.audit_id)
        forms_rows = [
            [f.get("page_url"), f.get("action_url"), f.get("method"), f.get("form_id"), f.get("has_password"), f.get("is_insecure")]
            for f in forms
        ]
        add_sheet("Forms", ["Page URL", "Action URL", "Method", "Form ID", "Has Password", "Is Insecure"], forms_rows)

        file_path = self.output_dir / "report.xlsx"
        wb.save(file_path)
        return str(file_path.absolute())

    async def generate_docx(self) -> str:
        data = await self._fetch_data()
        domain = data['audit'].get('domain', 'Unknown Domain')
        date_str = datetime.now().strftime("%Y-%m-%d")
        health_score = data['audit'].get('health_score', 0)
        
        doc = docx.Document()
        
        # Styles
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            
            # Header
            header = section.header
            header_para = header.paragraphs[0]
            header_para.text = f"SEO Audit Report - {domain}"
            header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Title Page
        doc.add_heading('SEO Audit Report', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Domain: {domain}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Date: {date_str}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        
        # TOC Placeholder
        doc.add_heading('Table of Contents', level=1)
        doc.add_paragraph('Update Table of Contents in Word using References > Update Table')
        doc.add_page_break()
        
        # Executive Summary
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph(f"Health Score: {health_score:.1f}/100", style='Intense Quote')
        
        critical_count = len([i for i in data['issues'] if i.get('severity') == 'critical'])
        warning_count = len([i for i in data['issues'] if i.get('severity') == 'warning'])
        info_count = len([i for i in data['issues'] if i.get('severity') == 'info'])
        
        doc.add_paragraph(f"Total Pages Crawled: {len(data['pages'])}")
        doc.add_paragraph(f"Critical Issues: {critical_count}")
        doc.add_paragraph(f"Warnings: {warning_count}")
        doc.add_paragraph(f"Info: {info_count}")
        
        doc.add_heading('Priority Fixes', level=2)
        top_issues = [i for i in data['issues'] if i.get('severity') == 'critical'][:5]
        for issue in top_issues:
            doc.add_paragraph(f"{issue.get('issue_type')} on {issue.get('url')}: {issue.get('message')}", style='List Bullet')
            
        doc.add_page_break()
        
        # All Issues Table
        doc.add_heading('All Issues', level=1)
        table = doc.add_table(rows=1, cols=4, style='Table Grid')
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Severity'
        hdr_cells[1].text = 'Type'
        hdr_cells[2].text = 'Message'
        hdr_cells[3].text = 'URL'
        
        for issue in sorted(data['issues'], key=lambda x: {'critical': 0, 'warning': 1, 'info': 2}.get(x.get('severity', 'info'), 3)):
            row_cells = table.add_row().cells
            row_cells[0].text = str(issue.get('severity')).upper()
            row_cells[1].text = str(issue.get('issue_type'))
            row_cells[2].text = str(issue.get('message'))
            row_cells[3].text = str(issue.get('url'))
            
        doc.add_page_break()
        
        # All Pages Table
        doc.add_heading('Appendix: All Pages', level=1)
        page_table = doc.add_table(rows=1, cols=3, style='Table Grid')
        hdr_cells = page_table.rows[0].cells
        hdr_cells[0].text = 'URL'
        hdr_cells[1].text = 'Status Code'
        hdr_cells[2].text = 'Word Count'
        
        for page in data['pages']:
            row_cells = page_table.add_row().cells
            row_cells[0].text = str(page.get('url'))
            row_cells[1].text = str(page.get('status_code'))
            row_cells[2].text = str(page.get('word_count'))
            
        file_path = self.output_dir / "report.docx"
        doc.save(file_path)
        return str(file_path.absolute())

    async def generate_all(self) -> dict:
        results = {
            "html": await self.generate_html(),
            "csv": await self.generate_csv(),
            "excel": await self.generate_excel(),
            "docx": await self.generate_docx()
        }
        try:
            from .visualizer import LinkGraphVisualizer
            vis = LinkGraphVisualizer(self.db, self.audit_id)
            results["link_graph_html"] = await vis.export_d3_html(str(self.output_dir / "link_graph.html"))
            results["link_graph_gexf"] = await vis.export_gexf(str(self.output_dir / "link_graph.gexf"))
        except Exception as e:
            logger.warning(f"Link graph export error: {e}")
        return results
